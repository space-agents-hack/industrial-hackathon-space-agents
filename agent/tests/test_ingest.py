"""Case ingestion, one reader at a time, on files built here. No challenge data: the fixtures are
made in the test so nothing confidential is committed (AGENTS.md)."""

import io
import struct

from docx import Document
from openpyxl import Workbook
from PIL import Image
from pydantic_ai import BinaryContent

from agent_service.domain import ingest
from agent_service.domain.ingest import CaseFile, Limits, format_of, read_case

LONG = "The rear lamp housing is sealed to IP6K9K and uses red PMMA for the outer lens. " * 4


def png(width: int, height: int) -> bytes:
    out = io.BytesIO()
    Image.new("RGB", (width, height), "white").save(out, format="PNG")
    return out.getvalue()


def pdf(pages: list[str]) -> bytes:
    """A minimal PDF with one text line per page; an empty string makes a page with no text."""
    objects: list[bytes] = []
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(len(pages)))
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    font = 3 + 2 * len(pages)
    for i, text in enumerate(pages):
        stream = f"BT /F1 10 Tf 20 700 Td ({text}) Tj ET".encode() if text else b""
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {4 + 2 * i} 0 R "
            f"/Resources << /Font << /F1 {font} 0 R >> >> >>".encode()
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    )
    return out.getvalue()


def texts(parts: list[ingest.Part]) -> str:
    return "\n".join(p for p in parts if isinstance(p, str))


def images(parts: list[ingest.Part]) -> list[BinaryContent]:
    return [p for p in parts if isinstance(p, BinaryContent)]


def test_extensions_map_to_the_formats_the_database_accepts() -> None:
    assert format_of("Spec.MD") == "md"
    assert format_of("view.JPG") == "jpeg"
    assert format_of("housing.stp") == "step"
    assert format_of("housing.igs") == "iges"
    assert format_of("notes.doc") is None
    assert format_of("README") is None


def test_text_is_read_as_it_is_and_labelled_with_its_file() -> None:
    content = read_case([CaseFile("spec.md", "md", b"# Rear lamp\nM1")])

    assert "## Case file: spec.md (md)" in texts(content.parts)
    assert "# Rear lamp\nM1" in texts(content.parts)
    [report] = content.reports
    assert report.characters == len("# Rear lamp\nM1")
    assert report.error is None


def test_text_that_is_not_utf8_is_read_and_the_trace_says_so() -> None:
    content = read_case([CaseFile("old.txt", "txt", "Größe".encode("latin-1"))])

    assert "Größe" in texts(content.parts)
    assert "Latin-1" in content.reports[0].notes[0]


def test_a_large_image_is_downscaled_and_a_small_one_sent_unchanged() -> None:
    small = png(800, 600)
    content = read_case(
        [CaseFile("view.png", "png", small), CaseFile("big.png", "png", png(4000, 2000))]
    )

    first, second = images(content.parts)
    assert first.data == small
    with Image.open(io.BytesIO(second.data)) as img:
        assert max(img.size) == ingest.MAX_IMAGE_SIDE
    assert second.identifier == "big.png"
    assert [r.images for r in content.reports] == [1, 1]


def test_pdf_text_keeps_its_page_and_a_page_without_text_goes_as_an_image() -> None:
    content = read_case([CaseFile("datasheet.pdf", "pdf", pdf([LONG, ""]))])

    assert "[page 1]\nThe rear lamp housing" in texts(content.parts)
    [drawing] = images(content.parts)
    assert drawing.identifier == "datasheet.pdf page 2"
    report = content.reports[0]
    assert report.pages == 2
    assert any("page(s) 2" in n for n in report.notes)


def test_word_keeps_headings_tables_and_embedded_images() -> None:
    doc = Document()
    doc.add_heading("Materials", level=2)
    doc.add_paragraph("Outer lens in PMMA.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Housing"
    table.rows[0].cells[1].text = "ASA/PC"
    doc.add_picture(io.BytesIO(png(200, 100)))
    out = io.BytesIO()
    doc.save(out)

    content = read_case([CaseFile("spec.docx", "docx", out.getvalue())])

    text = texts(content.parts)
    assert "## Materials" in text
    assert "Housing | ASA/PC" in text
    assert len(images(content.parts)) == 1


def test_each_spreadsheet_sheet_becomes_csv() -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "BOM"
    sheet.append(["Part", "Material"])
    sheet.append(["Lens", "PMMA"])
    out = io.BytesIO()
    book.save(out)

    content = read_case([CaseFile("bom.xlsx", "xlsx", out.getvalue())])

    assert "[sheet BOM]\nPart,Material\r\nLens,PMMA" in texts(content.parts)


STEP = b"""ISO-10303-21;
HEADER;
FILE_DESCRIPTION(('rear lamp assembly'),'2;1');
FILE_NAME('lamp.stp','2026-10-24',('someone'),(''),'','','');
ENDSEC;
DATA;
#1=PRODUCT('LAMP','Rear lamp','',(#90));
#2=PRODUCT('LENS','Outer lens','',(#90));
#11=PRODUCT_DEFINITION_FORMATION('','',#1);
#12=PRODUCT_DEFINITION_FORMATION('','',#2);
#21=PRODUCT_DEFINITION('design','',#11,#80);
#22=PRODUCT_DEFINITION('design','',#12,#80);
#31=NEXT_ASSEMBLY_USAGE_OCCURRENCE('1','lens','',#21,#22,$);
ENDSEC;
END-ISO-10303-21;
"""


def test_step_gives_header_parts_and_assembly_structure() -> None:
    content = read_case([CaseFile("lamp.stp", "step", STEP)])

    text = texts(content.parts)
    assert "rear lamp assembly" in text
    assert "- LAMP / Rear lamp" in text
    assert "- LAMP / Rear lamp > LENS / Outer lens (x1)" in text
    assert "geometry not interpreted" in content.reports[0].notes[0]


def test_iges_gives_its_start_and_global_sections() -> None:
    start = "Rear lamp housing exported from CAD".ljust(72) + "S      1"
    glob = "1H,,1H;,8Hlamp.igs,".ljust(72) + "G      1"
    content = read_case([CaseFile("lamp.igs", "iges", f"{start}\n{glob}\n".encode())])

    text = texts(content.parts)
    assert "Rear lamp housing exported from CAD" in text
    assert "8Hlamp.igs" in text


def test_stl_gives_triangles_and_bounding_box_in_both_encodings() -> None:
    ascii_stl = b"""solid lens
facet normal 0 0 1
outer loop
vertex 0 0 0
vertex 120 0 0
vertex 0 40 15
endloop
endfacet
endsolid lens
"""
    binary = b"lens".ljust(80, b"\0") + struct.pack("<I", 1)
    binary += struct.pack("<12fH", 0, 0, 1, 0, 0, 0, 120, 0, 0, 0, 40, 15, 0)

    content = read_case([CaseFile("a.stl", "stl", ascii_stl), CaseFile("b.stl", "stl", binary)])

    text = texts(content.parts)
    assert text.count("Triangles: 1") == 2
    assert text.count("Bounding box: 120.0 x 40.0 x 15.0") == 2


def test_a_broken_file_is_reported_and_the_others_still_read() -> None:
    content = read_case(
        [CaseFile("broken.pdf", "pdf", b"not a pdf"), CaseFile("spec.md", "md", b"M1")]
    )

    broken, spec = content.reports
    assert broken.error is not None and broken.error.startswith("Could not be read")
    assert "could not be read" in texts(content.parts)
    assert spec.error is None and spec.characters == 2


def test_past_the_limits_the_rest_is_reported_not_silently_dropped() -> None:
    content = read_case(
        [
            CaseFile("a.md", "md", b"x" * 30),
            CaseFile("b.md", "md", b"y" * 30),
            CaseFile("c.png", "png", png(10, 10)),
            CaseFile("d.png", "png", png(10, 10)),
        ],
        Limits(max_text_chars=40, max_images=1),
    )

    a, b, c, d = content.reports
    assert a.characters == 20 and b.characters == 20
    assert all(any("cut" in n for n in r.notes) for r in (a, b))
    assert c.images == 1
    assert d.images == 0 and any("Not sent" in n for n in d.notes)


def test_the_first_message_starts_with_the_note_and_keeps_upload_order() -> None:
    content = read_case([CaseFile("spec.md", "md", b"M1"), CaseFile("v.png", "png", png(5, 5))])

    message = ingest.first_message("Find the requirements.", content)

    assert isinstance(message[0], str) and message[0].startswith("Find the requirements.")
    assert "2 case file(s)" in message[0]
    assert "spec.md" in str(message[1])
    assert isinstance(message[-1], BinaryContent) and message[-1].identifier == "v.png"


def test_a_large_file_first_does_not_crowd_out_the_specification() -> None:
    content = read_case(
        [
            CaseFile("bom.csv", "csv", b"part,qty\n" * 1000),
            CaseFile("spec.md", "md", b"Rear lamp, category M1."),
            CaseFile("drawing.txt", "txt", b"z" * 5000),
        ],
        Limits(max_text_chars=3000, max_images=0),
    )

    bom, spec, drawing = content.reports
    assert spec.characters == len("Rear lamp, category M1.")
    # The two large files split what the spec left, to within one character.
    assert bom.characters + drawing.characters == 3000 - spec.characters
    assert abs(bom.characters - drawing.characters) <= 1


def test_only_the_pdf_pages_that_fit_the_image_limit_are_rendered(monkeypatch) -> None:
    rendered: list[int] = []
    real = ingest._render_pdf_page

    def counting(data: bytes, number: int) -> ingest.ImageData:
        rendered.append(number)
        return real(data, number)

    monkeypatch.setattr(ingest, "_render_pdf_page", counting)
    content = read_case([CaseFile("scan.pdf", "pdf", pdf([""] * 6))], Limits(max_images=2))

    assert rendered == [1, 2]
    assert content.reports[0].images == 2
    assert any("Not sent" in n and "page 6" in n for n in content.reports[0].notes)


def test_a_hash_sign_inside_a_step_name_is_not_read_as_a_reference() -> None:
    step = STEP.replace(b"'1','lens',''", b"'1','Lens #1 occurrence',''")
    assert step != STEP

    text = texts(read_case([CaseFile("lamp.stp", "step", step)]).parts)

    assert "- LAMP / Rear lamp > LENS / Outer lens (x1)" in text
