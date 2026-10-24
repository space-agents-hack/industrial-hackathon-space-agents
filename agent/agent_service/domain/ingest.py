"""Case ingestion: the files an engineer uploaded, turned into what the model reads (issue #8,
docs/se IF-1 and IF-2). Deterministic and model-free, so every reader is tested on a fixture.

The result is *case content*: text by file and page, plus images. It is deliberately not a fixed
"case profile": Certivity could not say which fields one should have, so the regulation step reads
the content itself, with file and page next to every block, and draws its own conclusions.

Every file produces a report (`FileReport`) that goes into the trace: what was read, and anything
skipped, truncated or unreadable. A file that fails to parse is reported and the run goes on with
the rest; nothing is dropped silently (SYS-2).

What this does NOT do:
- OCR. A PDF page with almost no text (a drawing, a scan) is rendered and sent as an image, which
  needs a vision model.
- Interpret CAD geometry. STEP and IGES are read for their header, part names and assembly
  structure; STL for its name, triangle count and bounding box.
- Read .doc, .ppt, .pptx or ZIP. The database refuses those formats (run_files.format).

Later: PDFs should go through the regulation ingestion layer (#3) once it exists, so that case
PDFs and regulation PDFs are split into pages the same way.
"""

import csv
import io
import re
import struct
from collections.abc import Sequence
from dataclasses import dataclass, field
from typing import Literal, get_args

from pydantic_ai import BinaryContent

Format = Literal[
    "md", "txt", "csv", "json", "png", "jpeg", "webp", "pdf", "docx", "xlsx", "step", "iges", "stl"
]
FORMATS: tuple[Format, ...] = get_args(Format)

# Extensions as people type them, to the format the database stores (run_files.format).
EXTENSIONS: dict[str, Format] = {
    **{f: f for f in FORMATS},
    "markdown": "md",
    "jpg": "jpeg",
    "stp": "step",
    "igs": "iges",
}

# The long side the providers recommend; larger images are downscaled by the provider anyway and
# cost more tokens for nothing.
MAX_IMAGE_SIDE = 1568
# Below this many characters a PDF page is taken to be a drawing or a scan, and rendered.
MIN_PAGE_TEXT = 200


@dataclass(frozen=True)
class Limits:
    """What one run sends the model. Above this the rest is reported as skipped, not sent.

    150,000 characters is about 40,000 tokens: room left in a 200k context for regulations.
    The text is shared fairly between files (`_fair_shares`); images go in upload order.
    """

    max_text_chars: int = 150_000
    max_images: int = 20


DEFAULT_LIMITS = Limits()


@dataclass(frozen=True)
class CaseFile:
    filename: str
    format: Format
    data: bytes
    # What the caller already knows about the file (a checksum mismatch, say), for its report.
    notes: tuple[str, ...] = ()


@dataclass
class FileReport:
    filename: str
    # str, not Format: a file whose format this agent cannot read is still reported as it came.
    format: str
    bytes: int
    pages: int | None = None
    characters: int = 0
    images: int = 0
    notes: list[str] = field(default_factory=list)
    error: str | None = None

    def payload(self) -> dict[str, object]:
        return {
            "filename": self.filename,
            "format": self.format,
            "bytes": self.bytes,
            "pages": self.pages,
            "characters": self.characters,
            "images": self.images,
            "notes": self.notes,
            "error": self.error,
        }


Part = str | BinaryContent


@dataclass
class CaseContent:
    parts: list[Part]
    reports: list[FileReport]


def format_of(filename: str) -> Format | None:
    _, dot, ext = filename.rpartition(".")
    return EXTENSIONS.get(ext.lower()) if dot else None


# ------------------------------------------------------------------------------------------
# Reading one file. Each reader returns blocks in order: text or an image, with the page it
# came from where the format has pages, and the parser that made it. This is the standard shape
# stored in case_file_blocks, so a block reads the same whether it was parsed now or last week.
# ------------------------------------------------------------------------------------------


@dataclass
class ImageData:
    data: bytes
    media_type: str


@dataclass
class PdfPage:
    """A PDF page to render only when it will be used. Rendering every drawing page up front cost
    the time and memory of pages that were then not sent: a scanned 60-page PDF rendered 60
    images to send 20 (found by the code review thread)."""

    data: bytes
    number: int


@dataclass
class Block:
    content: str | ImageData | PdfPage
    label: str = ""
    page: int | None = None
    # "pypdf", "pdfium-render", "model:<model>"... A citation of a model's transcription has to
    # be checked against the source, so the parser travels with the block.
    parser: str = ""


_Block = Block


def _text(data: bytes, report: FileReport) -> list[_Block]:
    try:
        return [Block(data.decode("utf-8-sig"), parser="text")]
    except UnicodeDecodeError:
        report.notes.append("Not valid UTF-8; read as Latin-1, so some characters may be wrong.")
        return [Block(data.decode("latin-1"), parser="text")]


def _image(data: bytes) -> ImageData:
    from PIL import Image

    with Image.open(io.BytesIO(data)) as img:
        img.load()
        if max(img.size) <= MAX_IMAGE_SIDE and img.format in ("PNG", "JPEG", "WEBP"):
            return ImageData(data, f"image/{img.format.lower()}")
        photo = img.format == "JPEG"
        img.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE))
        out = io.BytesIO()
        if photo:
            img.convert("RGB").save(out, format="JPEG", quality=85)
            return ImageData(out.getvalue(), "image/jpeg")
        # PNG keeps line drawings sharp; exploded views are line drawings more often than photos.
        if img.mode not in ("RGB", "RGBA", "L", "LA"):
            img = img.convert("RGBA")
        img.save(out, format="PNG", optimize=True)
        return ImageData(out.getvalue(), "image/png")


def _pdf(data: bytes, report: FileReport) -> list[_Block]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    report.pages = len(reader.pages)
    blocks: list[_Block] = []
    drawn: list[int] = []
    for number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            blocks.append(Block(text, f"page {number}", number, "pypdf"))
        if len(text) < MIN_PAGE_TEXT:
            drawn.append(number)
            blocks.append(Block(PdfPage(data, number), f"page {number}", number, "pdfium-render"))
    if drawn:
        report.notes.append(f"Little or no text on page(s) {_ranges(drawn)}; read as images.")
    return blocks


def _render_pdf_page(data: bytes, number: int) -> ImageData:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(data)
    try:
        page = pdf[number - 1]
        width, height = page.get_size()
        scale = MAX_IMAGE_SIDE / max(width, height, 1)
        bitmap = page.render(scale=scale)
        out = io.BytesIO()
        bitmap.to_pil().save(out, format="PNG", optimize=True)
        return ImageData(out.getvalue(), "image/png")
    finally:
        pdf.close()


def _docx(data: bytes, report: FileReport) -> list[_Block]:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(data))
    lines: list[str] = []
    # Body order, so a table stays next to the paragraph that introduces it.
    for item in doc.iter_inner_content():
        if isinstance(item, Paragraph):
            if item.text.strip():
                style = (item.style.name if item.style is not None else None) or ""
                heading = re.fullmatch(r"Heading ([1-6])", style)
                prefix = "#" * int(heading.group(1)) + " " if heading else ""
                lines.append(prefix + item.text)
        elif isinstance(item, Table):
            for row in item.rows:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells))
            lines.append("")
    blocks: list[_Block] = [Block("\n".join(lines), parser="python-docx")]
    report.notes.append("Word files have no fixed pages; cite headings or paragraphs instead.")
    for rel in doc.part.rels.values():
        if "image" in rel.reltype and not rel.is_external:
            try:
                image = _image(rel.target_part.blob)
                blocks.append(Block(image, "embedded image", parser="python-docx"))
            except Exception:
                report.notes.append("An embedded image could not be read and was skipped.")
    return blocks


def _xlsx(data: bytes, report: FileReport) -> list[_Block]:
    from openpyxl import load_workbook

    # data_only: the values Excel last computed, not the formulas.
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    blocks: list[_Block] = []
    try:
        for sheet in book.worksheets:
            out = io.StringIO()
            writer = csv.writer(out)
            for row in sheet.iter_rows(values_only=True):
                if any(v is not None for v in row):
                    writer.writerow(["" if v is None else v for v in row])
            blocks.append(Block(out.getvalue(), f"sheet {sheet.title}", parser="openpyxl"))
    finally:
        book.close()
    report.pages = len(blocks)
    return blocks


_STEP_ENTITY = re.compile(r"#(\d+)\s*=\s*([A-Z0-9_]+)\s*\((.*?)\)\s*;", re.S)
_STEP_STRING = re.compile(r"'((?:[^']|'')*)'")
_STEP_REF = re.compile(r"#(\d+)")


def _step(data: bytes, report: FileReport) -> list[_Block]:
    """Header, products and assembly structure of a STEP file (ISO 10303-21).

    A regex over entities, not a STEP parser: good enough for names and structure, which is all
    the requirements step can use. Geometry entities are counted, not read.
    """
    text = data.decode("latin-1")
    header = text.split("ENDSEC;", 1)[0] if "HEADER;" in text else ""
    entities = {m.group(1): (m.group(2), m.group(3)) for m in _STEP_ENTITY.finditer(text)}

    def strings(args: str) -> list[str]:
        return [s.replace("''", "'") for s in _STEP_STRING.findall(args)]

    def refs(args: str) -> list[str]:
        # Strings first: a name such as 'Lens #1' would otherwise read as a reference to entity 1.
        return _STEP_REF.findall(_STEP_STRING.sub("''", args))

    product_name: dict[str, str] = {}
    for eid, (kind, args) in entities.items():
        if kind == "PRODUCT":
            s = strings(args)
            product_name[eid] = " / ".join(x for x in s[:3] if x) or f"#{eid}"

    def definition_name(pd_id: str) -> str:
        # PRODUCT_DEFINITION -> PRODUCT_DEFINITION_FORMATION(...) -> PRODUCT
        _, pd_args = entities.get(pd_id, ("", ""))
        for formation in refs(pd_args):
            kind, f_args = entities.get(formation, ("", ""))
            if kind.startswith("PRODUCT_DEFINITION_FORMATION"):
                for product in refs(f_args):
                    if product in product_name:
                        return product_name[product]
        return f"#{pd_id}"

    links: dict[tuple[str, str], int] = {}
    for kind, args in entities.values():
        if kind == "NEXT_ASSEMBLY_USAGE_OCCURRENCE":
            r = refs(args)
            if len(r) >= 2:
                key = (definition_name(r[0]), definition_name(r[1]))
                links[key] = links.get(key, 0) + 1

    lines = ["Header:", header.replace("HEADER;", "").strip(), ""]
    lines.append(f"Products ({len(product_name)}):")
    lines += [f"- {name}" for name in sorted(set(product_name.values()))]
    if links:
        lines += ["", "Assembly structure (parent > child, occurrences):"]
        lines += [f"- {p} > {c} (x{n})" for (p, c), n in sorted(links.items())]
    lines += ["", f"{len(entities)} entities in total; geometry not interpreted."]
    report.notes.append("CAD: names and assembly structure read; geometry not interpreted.")
    return [Block("\n".join(lines), parser="step-header")]


def _iges(data: bytes, report: FileReport) -> list[_Block]:
    """Start and Global sections of an IGES file: the description, file name, author and units."""
    text = data.decode("latin-1")
    start, glob = [], []
    for line in text.splitlines():
        if len(line) >= 73:
            section = line[72]
            if section == "S":
                start.append(line[:72].rstrip())
            elif section == "G":
                glob.append(line[:72].rstrip())
    body = "Start section:\n" + "\n".join(start) + "\n\nGlobal section:\n" + "".join(glob)
    report.notes.append("CAD: header sections read; geometry not interpreted.")
    return [Block(body, parser="iges-header")]


def _stl(data: bytes, report: FileReport) -> list[_Block]:
    """Name, triangle count and bounding box. STL carries no units; they are usually mm."""
    points: list[tuple[float, float, float]] = []
    name = ""
    binary_count = struct.unpack_from("<I", data, 80)[0] if len(data) >= 84 else -1
    if len(data) == 84 + 50 * binary_count:
        name = data[:80].rstrip(b"\0 ").decode("latin-1", "replace")
        for i in range(binary_count):
            v = struct.unpack_from("<9f", data, 84 + 50 * i + 12)
            points += [v[0:3], v[3:6], v[6:9]]
        triangles = binary_count
    else:
        text = data.decode("latin-1")
        first = text.split("\n", 1)[0]
        name = first.removeprefix("solid").strip()
        for m in re.finditer(r"vertex\s+(\S+)\s+(\S+)\s+(\S+)", text):
            points.append((float(m.group(1)), float(m.group(2)), float(m.group(3))))
        triangles = len(points) // 3
    lines = [f"Name: {name or '(none)'}", f"Triangles: {triangles}"]
    if points:
        lo = [min(p[i] for p in points) for i in range(3)]
        hi = [max(p[i] for p in points) for i in range(3)]
        size = " x ".join(f"{h - low:.1f}" for low, h in zip(lo, hi, strict=True))
        lines.append(f"Bounding box: {size} (units not stated in STL; usually mm)")
    report.notes.append("CAD: mesh size read; geometry not interpreted.")
    return [Block("\n".join(lines), parser="stl-mesh")]


def _read(file: CaseFile, report: FileReport) -> list[_Block]:
    match file.format:
        case "md" | "txt" | "csv" | "json":
            return _text(file.data, report)
        case "png" | "jpeg" | "webp":
            return [Block(_image(file.data), parser="pillow")]
        case "pdf":
            return _pdf(file.data, report)
        case "docx":
            return _docx(file.data, report)
        case "xlsx":
            return _xlsx(file.data, report)
        case "step":
            return _step(file.data, report)
        case "iges":
            return _iges(file.data, report)
        case "stl":
            return _stl(file.data, report)


# ------------------------------------------------------------------------------------------
# All files of a run, within the limits.
# ------------------------------------------------------------------------------------------


@dataclass
class ParsedFile:
    filename: str
    format: str
    report: FileReport
    blocks: list[Block]


def parse(file: CaseFile) -> ParsedFile:
    """One file to blocks. Never raises: a file that cannot be read has an error in its report
    and no blocks, and the caller decides whether a model gets a second try."""
    report = FileReport(file.filename, file.format, len(file.data), notes=list(file.notes))
    try:
        blocks = _read(file, report)
    except Exception as exc:
        report.error = f"Could not be read: {type(exc).__name__}: {exc}"[:500]
        blocks = []
    return ParsedFile(file.filename, file.format, report, blocks)


def read_case(files: Sequence[CaseFile], limits: Limits = DEFAULT_LIMITS) -> CaseContent:
    return assemble([parse(f) for f in files], limits)


def assemble(files: Sequence[ParsedFile], limits: Limits = DEFAULT_LIMITS) -> CaseContent:
    """Blocks of every file into the model's message parts, within the limits. The report's
    counts say what was sent; what was not is in its notes."""
    parts: list[Part] = []
    reports: list[FileReport] = []
    demands = [sum(len(b.content) for b in f.blocks if isinstance(b.content, str)) for f in files]
    shares = _fair_shares(demands, limits.max_text_chars)
    images_left = limits.max_images

    for parsed, chars_left in zip(files, shares, strict=True):
        report = parsed.report
        report.characters = report.images = 0
        reports.append(report)
        if report.error is not None and not parsed.blocks:
            parts.append(f"## Case file: {parsed.filename}\n(This file could not be read.)")
            continue

        texts: list[str] = [f"## Case file: {parsed.filename} ({parsed.format})"]
        skipped: list[str] = []
        for block in parsed.blocks:
            label, content = block.label, block.content
            by = (
                " (transcribed by a model; check against the source)"
                if block.parser.startswith("model:")
                else ""
            )
            where = f"[{label}{by}]\n" if label or by else ""
            if isinstance(content, str):
                if chars_left <= 0:
                    skipped.append(label or "text")
                    continue
                if len(content) > chars_left:
                    content = content[:chars_left]
                    report.notes.append(f"Text cut at {label or 'the end'}: run text limit.")
                chars_left -= len(content)
                report.characters += len(content)
                texts.append(where + content)
            else:
                if images_left <= 0:
                    skipped.append(f"image ({label})" if label else "image")
                    continue
                if isinstance(content, PdfPage):
                    try:
                        content = render(content)
                    except Exception as exc:
                        report.notes.append(
                            f"{label} could not be rendered: {type(exc).__name__}: {exc}"[:300]
                        )
                        continue
                images_left -= 1
                report.images += 1
                texts.append(f"{where}(image follows)")
                parts.append("\n\n".join(texts))
                texts = []
                parts.append(
                    BinaryContent(
                        content.data,
                        media_type=content.media_type,
                        identifier=f"{parsed.filename} {label}".strip(),
                    )
                )
        if skipped:
            report.notes.append(f"Not sent (run limits): {', '.join(skipped)}.")
            texts.append(f"(Not sent because of the run's limits: {', '.join(skipped)}.)")
        if texts:
            parts.append("\n\n".join(texts))

    return CaseContent(parts, reports)


def render(page: PdfPage) -> ImageData:
    return _render_pdf_page(page.data, page.number)


def _fair_shares(demands: Sequence[int], budget: int) -> list[int]:
    """Split the text budget so a small file is sent whole and the large ones share the rest.

    First come, first served let one large spreadsheet ahead of the specification use the whole
    budget, and the specification was not sent at all. Max-min fair: each file gets what it needs
    or an equal part of what is left, whichever is smaller.
    """
    shares = [0] * len(demands)
    left, waiting = budget, len(demands)
    for i in sorted(range(len(demands)), key=lambda i: demands[i]):
        shares[i] = min(demands[i], left // waiting)
        left -= shares[i]
        waiting -= 1
    return shares


def first_message(task: str, content: CaseContent) -> list[Part]:
    """The opening user message: the engineer's note, then every case file in upload order."""
    intro = (
        f"{task}\n\nThe engineer attached {len(content.reports)} case file(s). Their content "
        "follows, each block labelled with its file and, where the format has them, its page. "
        "Cite file and page for anything you take from them."
    )
    return [intro, *content.parts]


def _ranges(numbers: list[int]) -> str:
    out: list[str] = []
    start = prev = numbers[0]
    for n in [*numbers[1:], None]:
        if n is not None and n == prev + 1:
            prev = n
            continue
        out.append(f"{start}" if start == prev else f"{start}-{prev}")
        if n is not None:
            start = prev = n
    return ", ".join(out)
